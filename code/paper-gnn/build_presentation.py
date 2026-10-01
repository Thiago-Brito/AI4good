"""Build editable slides and a PDF companion from verified local evidence.

Run with the project virtual environment after capture_presentation.py.
PDF is rendered from the companion HTML; PPTX uses editable native text.
"""
from pathlib import Path
import hashlib
import html
import json

import fitz
from PIL import Image, ImageOps, ImageDraw
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'outputs/apresentacao-ordered-gnn'
PROJECT = ROOT / 'code/paper-gnn'
PAPER_URL = 'https://openreview.net/forum?id=wKPmPBHSnT6'
summary = json.loads((PROJECT / 'outputs/main/summary.json').read_text())
slides = []


def add(title, body, notes, image=None, code=None, source='Código e evidências locais • code/paper-gnn'):
    slides.append(dict(title=title, body=body, notes=notes, image=image, code=code, source=source))


def pct(value):
    return f'{100 * value:.2f}'.replace('.', ',')



def compact_slides():
    """Keep the five-slide narrative requested for the actual presentation."""
    slides.clear()
    add('O artigo: Ordered GNN • ICLR 2023', [
        'Song, Zhou, Wang e Lin: uma rede neural que aprende com itens e suas ligações.',
        'Ideia: controlar a troca de informações entre vizinhos para preservar suas diferenças.',
        'Nosso objetivo: executar a referência, modificar a arquitetura e comparar.'
    ], 'O título completo é Ordered GNN: Ordering Message Passing to Deal with Heterophily and Over-smoothing. Mostre a primeira página. GNN significa rede neural para grafos: itens ligados entre si. Aqui os itens são documentos e as ligações são citações. Nosso protocolo é compacto, com adaptações para CPU, largura e orçamento diferentes; não é uma reprodução integral das tabelas do artigo.', image='artigo-pagina1.png', source=PAPER_URL)
    add('A base Cora: documentos ligados por citações', [
        '2.708 documentos, 7 assuntos e 1.433 indicadores de palavras por documento.',
        'Citação = um documento menciona outro nas referências. Isso cria uma ligação.',
        'Classe = assunto. Exemplo: classe 3 significa “Redes neurais”.',
        'Treino: 1.208 • validação: 500 • teste: 1.000 documentos.'
    ], 'Base de dados é o conjunto de exemplos. Rede de referência é o modelo com que comparamos nossa versão. data.x guarda indicadores de palavras, data.edge_index as ligações e data.y os assuntos corretos. Não usamos PDFs completos como entrada. O grafo tem 5.278 ligações usadas nos dois sentidos. Todos os atributos e ligações ficam visíveis, mas apenas os rótulos de treino entram no gradiente. Validação escolhe os pesos salvos; teste mede o resultado final.', code="dataset = Planetoid(\n    str(OUTPUTS / 'dataset'),\n    'Cora', split='full',\n    transform=NormalizeFeatures()\n)\ndata = dataset[0]", source='data.py • outputs/dataset/manifest.json • verify_categories.py')
    add('O código aprende; a aplicação faz a previsão', [
        'A GNN resume as palavras e faz quatro rodadas de troca entre documentos ligados.',
        'Treinar ajusta os números internos da rede para reduzir os erros.',
        'O treino já foi feito: a aplicação carrega os pesos e calcula novas previsões.',
        'Print real, documento 1713: assunto Teoria. A referência errou; a modificada acertou.'
    ], 'Cada documento vira uma representação com 64 números. Cada rodada agrega informações dos vizinhos e controla a mistura com gates. Na referência, a última rodada alimenta a classificação. Treinamos cada arquitetura do zero nas sementes 42, 43 e 44: seis execuções. initial.pt, model.pt e history.csv registram pesos iniciais, finais e evolução. audit.py verifica dados, mudanças nos pesos e métricas recalculadas. Na aplicação não ocorre novo treino. Abra localhost:8765 e escolha 1713: a referência prevê Aprendizado baseado em casos; a modificada prevê Teoria. As barras são probabilidades, não garantias de acerto.', image='aplicacao-documento-1713.png', source='train.py • infer.py • captura real da aplicação • semente 42')
    add('Minha modificação: combinar mais informações', [
        'Referência: decide usando a representação da quarta e última rodada.',
        'Mudança 1: cada rodada aprende seu próprio controle da troca de informações.',
        'Mudança 2: uma pequena rede transforma os dados e soma sua saída à entrada do bloco.',
        'Mudança 3: aprende a combinar a entrada e as quatro rodadas antes de decidir.',
        '99.223 → 174.344 parâmetros. A versão modificada foi treinada do zero.'
    ], 'Mostre MultiScaleOrderedGNN em models.py. A primeira mudança usa gates independentes. A segunda são blocos residuais: transformação 64 para 128 para 64, GELU e normalização. A terceira usa states, depth_score e softmax para aprender cinco pesos por documento e combinar profundidades. Parâmetros são números ajustados no treino: o aumento foi de 75,71%. As modificações são deste projeto, não dos autores. Foram avaliadas juntas, sem isolar o efeito de cada uma.', source='models.py • MultiScaleOrderedGNN • outputs/main/')
    add('Resultados e relatório: o que concluímos?', [
        f"Acertos no teste: referência {pct(summary['original']['accuracy']['mean'])}% → modificada {pct(summary['modified']['accuracy']['mean'])}%.",
        'Média de três sementes: a rede maior não melhorou o resultado geral.',
        'Há acertos e falhas: no documento 1708, as duas redes erram.',
        'O relatório reúne motivação, método, gráficos, métricas e exemplos.',
        'Escrita e revisão com apoio de IA; resultados conferidos por testes e auditoria.'
    ], 'A queda média foi de 0,13 ponto percentual no mesmo split. O desvio padrão da acurácia foi 0,78 ponto na referência e 0,26 na modificada. F1 macro: 86,19% contra 86,00%; considera as categorias com o mesmo peso. Não fizemos teste de significância. O caso 1713 não prova superioridade geral. No 1708, o esperado é Redes neurais e ambas dizem Teoria. Mostre o relatório de nove páginas no Overleaf ou /report.pdf: motivação, protocolo, mudanças, resultados quantitativos, qualitativos e limitações. As fontes estão em refs/ordered-gnn.md e seis modelos foram auditados. Apoio de IA não garante ausência absoluta de erros.', source='outputs/main/summary.json • outputs/audit.json • samplepaper.tex')


def build():
    compact_slides()
    OUT.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    navy, teal = '14304B', '008579'

    def text(slide, value, x, y, w, h, size=22, color=navy, bold=False, font='Arial'):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = shape.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = 0
        for i, line in enumerate(value.split('\n')):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.text = line
            p.font.name, p.font.size, p.font.bold = font, Pt(size), bold
            p.font.color.rgb = RGBColor.from_string(color)
            p.space_after = Pt(16)
        return shape

    sections = []
    for i, s in enumerate(slides, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string('F4F7FB')
        text(slide, 'ORDERED GNN / CORA / EXPERIMENTO LOCAL', .55, .28, 12, .35, 11, teal, True)
        text(slide, s['title'], .55, .88, 12.2, .8, 31, bold=True)
        visual = s['image'] or s['code']
        width = 5.3 if visual else 11.8
        text(slide, '\n'.join(s['body']), .6, 1.95, width, 4.85, 20 if visual else 23)
        if s['image']:
            path = OUT / s['image']
            iw, ih = Image.open(path).size
            scale = min(6.0 / iw, 4.8 / ih)
            w, h = iw * scale, ih * scale
            slide.shapes.add_picture(str(path), Inches(6.6 + (6-w)/2), Inches(1.95 + (4.8-h)/2), width=Inches(w), height=Inches(h))
        if s['code']:
            text(slide, s['code'], 6.45, 2.3, 6.25, 4.1, 17, teal, font='Consolas')
        text(slide, s['source'], .6, 7.04, 11.7, .25, 9)
        text(slide, f'{i:02}', 12.35, 6.98, .4, .35, 12, teal, True)
        slide.notes_slide.notes_text_frame.text = s['notes'] + '\n\nFonte: ' + s['source']
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width
            assert shape.top + shape.height <= prs.slide_height
        content = ''.join(f'<p>{html.escape(b)}</p>' for b in s['body'])
        right = f'<img src="{s["image"]}">' if s['image'] else f'<pre>{html.escape(s["code"])}</pre>' if s['code'] else ''
        sections.append(f'<section><header>ORDERED GNN / CORA / EXPERIMENTO LOCAL</header><h1>{html.escape(s["title"])}</h1><main class="{"two" if visual else "one"}"><article>{content}</article>{right}</main><footer>{html.escape(s["source"])}<b>{i:02}</b></footer></section>')
    prs.save(OUT / 'apresentacao-ordered-gnn.pptx')
    css = '''@page{size:1280px 720px;margin:0}*{box-sizing:border-box}body{margin:0;font-family:Arial;color:#14304b;background:#f4f7fb}section{width:1280px;height:720px;padding:30px 55px;position:relative;break-after:page}header{color:#008579;font-size:14px;font-weight:bold;letter-spacing:1px}h1{font-size:39px;margin:34px 0 32px}main{height:470px;display:grid;align-items:center;gap:45px}.two{grid-template-columns:1fr 1.12fr}.one{grid-template-columns:1fr}p{font-size:24px;line-height:1.27;margin:0 0 24px}.one p{font-size:28px}img{max-width:100%;max-height:465px;justify-self:center}pre{white-space:pre-wrap;font:21px/1.65 Consolas,monospace;color:#008579}footer{position:absolute;left:55px;right:55px;bottom:24px;font-size:12px}footer b{float:right;color:#008579}'''
    (OUT/'apresentacao.html').write_text('<!doctype html><html lang="pt-BR"><meta charset="utf-8"><title>Ordered GNN e Cora</title><style>'+css+'</style>'+''.join(sections)+'</html>', encoding='utf-8')
    (OUT/'roteiro-de-fala.md').write_text('# Roteiro de apresentação — cerca de 4 a 5 minutos\n\nCada slide tem notas de fala também no PowerPoint.\n\n'+ '\n\n'.join(f'## {i}. {s["title"]}\n\n{s["notes"]}\n\nFonte: {s["source"]}' for i,s in enumerate(slides,1)), encoding='utf-8')
    with sync_playwright() as p:
        browser = p.chromium.launch(channel='msedge', headless=True)
        page = browser.new_page(viewport={'width':1280,'height':720})
        page.goto((OUT/'apresentacao.html').as_uri(), wait_until='networkidle')
        overflow = page.evaluate('''() => [...document.querySelectorAll('section')].map((s,i)=>({slide:i+1, overflow:[...s.querySelectorAll('main p, main img, main pre')].some(e=>e.getBoundingClientRect().bottom>s.getBoundingClientRect().top+675)})).filter(x=>x.overflow)''')
        assert not overflow, overflow
        page.pdf(path=str(OUT/'apresentacao-ordered-gnn.pdf'), print_background=True, prefer_css_page_size=True)
        browser.close()
    pdf = fitz.open(OUT/'apresentacao-ordered-gnn.pdf')
    assert len(pdf) == len(slides)
    sheet = Image.new('RGB', (1280, ((len(slides) + 3) // 4)*205), '#dce5ed')
    for i, page in enumerate(pdf):
        pix = page.get_pixmap(matrix=fitz.Matrix(.5,.5))
        thumb = Image.frombytes('RGB', [pix.width,pix.height], pix.samples)
        thumb.thumbnail((310,175))
        x,y=(i%4)*320,(i//4)*205
        sheet.paste(thumb,(x,y))
        ImageDraw.Draw(sheet).text((x+5,y+178),f'Slide {i+1}',fill='black')
    sheet.save(OUT/'revisao-slides.png')
    manifest = {'slides':len(slides),'pdf_pages':len(pdf),'html_overflow':overflow,'pdf_render':'HTML companion; PPTX contains editable native text','sources':str(ROOT/'academy/papers/artigo-overleaf/refs/ordered-gnn.md'),'sha256':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in OUT.iterdir() if p.suffix in {'.png','.json','.pptx','.pdf'} and p.name!='manifest.json'}}
    (OUT/'manifest.json').write_text(json.dumps(manifest,indent=2,ensure_ascii=False),encoding='utf-8')
    print(f'{len(slides)} slides, PDF verified: {OUT}')


if __name__ == '__main__':
    build()
