"""Build an editable deck and a standalone SVG/HTML presentation."""
from pathlib import Path
import sys, html, json

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'outputs/slide-deps'))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement

OUT = ROOT / 'academy/papers/dsl-visual-ia/estudos/01-fundamentos/outputs/apresentacao'
OUT.mkdir(parents=True, exist_ok=True)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(16), Inches(9)
BG, FG, MUTED, ACC, CARD, GOLD = '#101D2D', '#F3F6FA', '#B5C5D6', '#58D9B3', '#1C3047', '#F3C979'
slides, notes = [], []
s = None
svg = []

def rect(x,y,w,h,fill=CARD):
    shape=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor.from_string(fill[1:]); shape.line.fill.background()
    svg.append(f'<rect x="{x*100}" y="{y*100}" width="{w*100}" height="{h*100}" rx="10" fill="{fill}"/>')

def txt(text,x,y,w=14,h=.6,size=24,color=FG,bold=False):
    box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame; tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name='Aptos'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color[1:])
        p.line_spacing=1.18; p.space_before=Pt(0);p.space_after=Pt(0)
        if line.startswith('https://'): p.runs[0].hyperlink.address=line
    for i,line in enumerate(text.split('\n')):
        element=f'<text xml:space="preserve" x="{x*100}" y="{y*100+size*1.17+i*size*1.65}" fill="{color}" font-size="{size*1.389}" font-weight="{700 if bold else 400}" font-family="Aptos,Segoe UI,sans-serif">{html.escape(line)}</text>'
        svg.append(f'<a href="{html.escape(line,quote=True)}">{element}</a>' if line.startswith('https://') else element)

def arrow(x,y,w=.55):
    shape=s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(x),Inches(y),Inches(w),Inches(.25))
    shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor.from_string(ACC[1:]); shape.line.fill.background()
    svg.append(f'<path d="M{x*100},{y*100+8} h{w*100-14} v-8 l14,12 -14,13 v-8 h{-w*100+14} Z" fill="{ACC}"/>')

def edge(x1,y1,x2,y2):
    line=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb=RGBColor.from_string(ACC[1:]);line.line.width=Pt(2)
    tip=OxmlElement('a:tailEnd');tip.set('type','triangle');line.line._get_or_add_ln().append(tip)
    svg.append(f'<defs><marker id="tip" markerWidth="8" markerHeight="8" refX="7" refY="4" orient="auto"><path d="M0,0 L8,4 L0,8 Z" fill="{ACC}"/></marker></defs><line x1="{x1*100}" y1="{y1*100}" x2="{x2*100}" y2="{y2*100}" stroke="{ACC}" stroke-width="3" marker-end="url(#tip)"/>')

def start(kicker,title,source='Exemplo didático próprio',note=''):
    global s,svg
    s=prs.slides.add_slide(prs.slide_layouts[6]); svg=[]
    s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor.from_string(BG[1:])
    rect(.65,.65,.12,.35,ACC); txt(kicker.upper(),.95,.65,size=13,color=ACC,bold=True)
    txt(title,.8,1.35,h=1.4,size=34,bold=True)
    txt(source,.8,8.45,w=13.7,h=.3,size=10,color=MUTED)
    txt(f'{len(slides)+1:02}',14.65,8.35,w=.6,size=16,color=ACC)
    s.notes_slide.notes_text_frame.text=note
    notes.append((title,note))

def end():
    slides.append('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1600 900" role="img" aria-label="'+html.escape(notes[-1][0],quote=True)+'"><rect width="1600" height="900" fill="'+BG+'"/>'+''.join(svg)+'</svg>')

def card(x,y,w,title,body):
    rect(x,y,w,2.45); txt(title,x+.25,y+.25,w-.5,size=22,color=ACC,bold=True)
    txt(body,x+.25,y+.95,w-.5,h=1.3,size=20,color=FG)

start('Fundamentos · palavras-chave do tema','DSL visual para orquestração de IA:\ncomeçando do básico', 'Base: GUIA-INICIAL.md · fontes S1–S4 no slide final', 'Apresente o objetivo: explicar as palavras-chave do tema e como elas se relacionam. O foco espacial é a organização dos blocos na tela. Duração sugerida: 10–12 minutos. Não é uma proposta de ferramenta para gerar slides.')
txt('Entendendo as palavras-chave e a organização\ndos blocos em uma interface visual.',.85,3.4,h=1.3,size=27,color=MUTED)
for x,t in [(1,'DESCREVER'),(6,'CONECTAR'),(11,'ORGANIZAR')]:
    rect(x,6,3.9,.9); txt(t,x+.3,6.23,w=3.4,size=22,color=ACC,bold=True)
arrow(5.2,6.33); arrow(10.2,6.33); end()

start('01 · Mapa do tema','Quais palavras precisamos entender?','S1–S4 · mapa conceitual didático', 'Leia o tema em partes. DSL indica uma linguagem especializada; visual indica como ela é representada; orquestração de IA indica que problema ela descreve. A organização espacial é um aspecto da representação visual que você quer estudar.')
card(.8,3.05,4.5,'DSL','Linguagem específica\npara um domínio.')
card(5.75,3.05,4.5,'VISUAL','Blocos, conexões\ne representação gráfica.')
card(10.7,3.05,4.5,'ORQUESTRAÇÃO DE IA','Coordenação de tarefas\nque utilizam IA.')
txt('Recorte de interesse: organização espacial dos blocos na tela.',.85,6.5,size=27,color=GOLD); end()

start('02 · Conceito','DSL: uma linguagem para um domínio','S1 · Fowler (2008)', 'DSL significa linguagem específica de domínio. Seu vocabulário é pensado para uma classe de problemas. Os blocos desta apresentação são inventados para explicar a ideia. Diferencie o domínio da linguagem de sua forma visual.')
txt('Domínio = a classe de problemas que queremos expressar.',.85,2.85,size=27,color=MUTED)
card(.8,4,6.95,'EXEMPLO CONHECIDO','SQL expressa consultas a dados.\nSeu vocabulário pertence a esse domínio.')
card(8.2,4,6.95,'NOSSO EXEMPLO DIDÁTICO','Pergunta, Buscar, MontarPrompt e LLM.\nCada bloco precisa de um significado.')
txt('Os blocos e as regras formam a linguagem.',.85,7.2,size=25,color=GOLD); end()

start('03 · Linguagem visual','Um desenho passa a ter regras','S1–S2 · Fowler · exemplo de linguagem proposto', 'Sintaxe é como montar expressões válidas; semântica é o que elas significam. Dê o exemplo: uma entrada que exige texto não recebe automaticamente uma imagem. As setas precisam declarar se carregam dados ou representam controle.')
card(.8,3.1,4.5,'VOCABULÁRIO','Quais blocos existem?\nEx.: Buscar e MontarPrompt.')
card(5.75,3.1,4.5,'CONEXÕES','O que pode ser ligado?\nEx.: texto → entrada textual.')
card(10.7,3.1,4.5,'SIGNIFICADO','O que cada bloco faz?\nEx.: recuperar trechos.')
txt('Uma seta precisa dizer o que transporta ou controla.',.85,6.6,size=28,color=GOLD); end()

start('04 · Orquestração de IA','Coordenar etapas, entradas e decisões','S3 · Anthropic (2024) · fluxo do guia, simplificado', 'Orquestrar é coordenar quem executa, com quais dados e em quais condições. LLM é um modelo de linguagem usado para gerar texto. O exemplo consulta uma base já preparada. A pergunta original e os trechos recuperados entram na montagem do prompt. Prompt é a entrada com instruções e contexto fornecida ao modelo. Busca e exibição não precisam ser executadas por IA. Um fluxo com LLM não é automaticamente um agente autônomo.')
for x,title,body in [(0.8,'Pergunta','texto recebido'),(3.75,'Buscar','recuperar trechos'),(6.7,'Montar prompt','pergunta + trechos'),(9.65,'LLM','gerar texto'),(12.6,'Mostrar','exibir resposta')]:
    rect(x,3.35,2.5,1.35); txt(title,x+.14,3.6,2.25,size=19,color=ACC,bold=True);txt(body,x+.14,4.13,2.25,size=13,color=MUTED)
    if x<12: arrow(x+2.58,3.94,.28)
edge(2.05,3.35,2.05,2.95);edge(2.05,2.95,7.95,2.95);edge(7.95,2.95,7.95,3.35)
txt('pergunta original',3.9,2.58,3.5,size=13,color=ACC)
txt('Workflow = fluxo de trabalho com etapas e relações definidas.',.85,5.7,size=26)
txt('Se a busca não encontrar trechos, o que deve acontecer?',.85,6.55,size=25,color=GOLD); end()

start('05 · Espacialidade','O espaço também faz parte da representação','S4 · Moody, Heymans e Matulevičius (2010) · recorte didático', 'Usamos espacialidade como organização de elementos no espaço visual. A literatura sobre notações visuais considera a posição uma variável de representação. Isso não comprova que nosso projeto será melhor: o efeito sobre a compreensão precisa ser avaliado.')
card(.8,3.1,4.5,'POSIÇÃO','Onde cada elemento fica?\nEx.: entrada à esquerda.')
card(5.75,3.1,4.5,'DISTÂNCIA','Quais blocos ficam próximos?\nEx.: etapas relacionadas.')
card(10.7,3.1,4.5,'AGRUPAMENTO','O que forma um conjunto?\nEx.: operações de busca.')
txt('Alinhamento e limites visuais completam essa organização.',.85,6.65,size=26,color=GOLD); end()

start('06 · Vocabulário dos diagramas','Blocos, portas e conexões','S2 e S4 · definições didáticas aplicadas ao fluxo', 'Nó é um elemento do grafo, representado aqui como bloco de operação. Porta é o ponto de entrada ou saída. Aresta é a ligação. Um tipo descreve a categoria do valor. No exemplo, Buscar recebe texto e produz uma lista de trechos. A conexão deve respeitar a entrada esperada pelo próximo bloco.')
rect(4.6,3.05,6.8,2.25);txt('BUSCAR DOCUMENTOS',5.3,3.75,5.5,size=28,color=ACC,bold=True)
rect(4.45,3.95,.3,.3,ACC);rect(11.25,3.95,.3,.3,ACC)
edge(1.3,4.1,4.4,4.1);edge(11.65,4.1,14.65,4.1)
txt('texto',1.5,3.45,2.8,size=21);txt('lista de trechos',11.85,3.45,3.2,size=20)
txt('porta de entrada',3.35,5.55,4,size=19,color=MUTED);txt('porta de saída',9.5,5.55,4,size=19,color=MUTED)
txt('Bloco = operação     Porta = entrada/saída     Conexão = relação',.85,6.85,size=25,color=GOLD); end()

start('07 · Organização espacial na prática','O mesmo fluxo, duas disposições na tela','S4 · motivação conceitual · diagramas próprios, sem comparação experimental', 'As duas disposições representam Entrada para Modelo para Saída. Na versão A, os blocos estão dispersos. Na B, estão alinhados na direção das setas. Os nomes e conexões são os mesmos. A hipótese é que a organização espacial ajude a acompanhar o fluxo; esse desenho não prova um ganho de compreensão.')
rect(.8,2.8,6.95,4.25);rect(8.2,2.8,6.95,4.25)
txt('A · BLOCOS DISPERSOS',1.1,3.1,6,size=19,color=ACC,bold=True);txt('B · BLOCOS ALINHADOS',8.5,3.1,6,size=19,color=ACC,bold=True)
edge(5.25,4.05,3.3,5.4);edge(3.3,5.65,5.6,6.35)
for x,y,label in [(5.15,3.75,'Entrada'),(1.45,5.2,'Modelo'),(5.25,6.05,'Saída'),(8.6,4.8,'Entrada'),(10.8,4.8,'Modelo'),(13,4.8,'Saída')]:
    rect(x,y,1.8,.65,'#2C4661');txt(label,x+.13,y+.15,1.6,size=18)
arrow(10.45,5,.25);arrow(12.65,5,.25)
txt('Mesmos blocos e conexões. A posição na tela é que mudou.',.85,7.5,size=25,color=GOLD);end()

start('08 · Distinção importante','Linguagem, editor e motor de execução','S2 · arquitetura ilustrativa do guia', 'A linguagem define os elementos, regras e significados. O editor é a interface usada para montar o fluxo; nele observamos o layout dos blocos. O motor executa as operações. No exemplo, mover um nó não altera sua execução, porque o comportamento é definido pelas conexões e regras. Algumas linguagens podem dar significado à posição: é uma decisão de projeto.')
card(.8,3.1,4.5,'LINGUAGEM','Define blocos, regras\ne significados.')
card(5.75,3.1,4.5,'EDITOR VISUAL','Permite montar, conectar\ne organizar os blocos.')
card(10.7,3.1,4.5,'MOTOR DE EXECUÇÃO','Executa as operações\ne transporta os dados.')
txt('Neste exemplo, mover um bloco não muda o comportamento.',.85,6.45,size=25,color=GOLD)
txt('O comportamento depende das conexões e das regras da linguagem.',.85,7.1,size=23,color=MUTED);end()

start('09 · Palavras-chave para pesquisa','Como nomear esse recorte?','S1–S4 · termos candidatos; sem contagem ou validação bibliométrica', 'Sugira organização espacial de blocos em interfaces visuais como expressão clara em português. Para pesquisar, combine termos de linguagem visual, layout de grafos e orquestração. As expressões são candidatas; ainda é necessário testar buscas e contagens nas bases da disciplina. Espacialidade sozinha é ampla.')
for i,(pt,en) in enumerate([('Linguagem específica de domínio','domain-specific language (DSL)'),('Programação visual','visual programming'),('Orquestração de fluxos de IA','AI workflow orchestration'),('Organização espacial de blocos','spatial layout / graph layout'),('Interface baseada em nós','node-based interface')]):
    y=2.85+i*.83;rect(.8,y,14.35,.7);txt(pt,1.05,y+.15,7,size=21);txt(en,8.5,y+.15,6.4,size=20,color=ACC)
txt('“Espacialidade” fica mais precisa quando ligada à interface e aos blocos.',.85,7.4,size=23,color=GOLD);end()

start('10 · O tema em uma frase','Como explicar o que estou estudando?',note='Leia a síntese como uma maneira de apresentar seu interesse atual, não como conclusão de pesquisa. A hipótese é que posição, alinhamento e agrupamento possam ajudar na compreensão. O público, a tarefa e a avaliação ainda precisam ser delimitados. Isso mantém o foco nas palavras-chave que o estudante quer explicar.')
txt('“Estudo uma linguagem visual para coordenar tarefas de IA,\ncom foco na organização dos blocos na tela\ne na compreensão do fluxo.”',.85,2.8,h=2,size=28,color=ACC,bold=True)
card(.8,5.45,6.95,'CONCEITOS CENTRAIS','DSL · linguagem visual · orquestração\nOrganização espacial · compreensão')
card(8.2,5.45,6.95,'PERGUNTA POSSÍVEL','Como a organização dos blocos ajuda\nas pessoas a compreender um fluxo?');end()

start('Referências','Base conceitual e limites', 'Registro completo: refs/FONTES-APRESENTACAO.md', 'As fontes sustentam os conceitos básicos. Diagramas e exemplos são didáticos. Não há demonstração empírica de que uma disposição seja melhor para os usuários deste projeto. O interesse é descrever os conceitos e delimitar a organização espacial dos blocos como recorte.')
refs=[('S1 · Fowler (2008) — Domain Specific Language','https://www.martinfowler.com/bliki/DomainSpecificLanguage.html'),('S2 · Fowler (2005) — Language Workbenches','https://martinfowler.com/articles/languageWorkbench.html'),('S3 · Anthropic (2024) — Building effective agents','https://www.anthropic.com/engineering/building-effective-agents'),('S4 · Moody, Heymans e Matulevičius (2010) — Visual syntax does matter','https://doi.org/10.1007/s00766-010-0100-1')]
for i,(label,url) in enumerate(refs):
    y=2.8+i*1.05;txt(label,.85,y,size=20,color=ACC,bold=True);txt(url,.85,y+.43,size=15,color=MUTED)
txt('Diagramas próprios · termos de busca candidatos · sem resultados experimentais',.85,7.55,size=20,color=GOLD);end()

prs.save(OUT/'dsl-visual-ia-espacialidade.pptx')
css='''*{box-sizing:border-box}body{margin:0;background:#080f19;color:#f3f6fa;font-family:Segoe UI,sans-serif}main{height:calc(100vh - 60px);display:grid;place-items:center}section{display:none;width:min(100vw,calc((100vh - 60px)*16/9))}section.active{display:block}svg{display:block;width:100%;height:auto}nav{height:60px;display:flex;align-items:center;justify-content:center;gap:16px}button{background:#1c3047;color:#f3f6fa;border:1px solid #59718b;border-radius:7px;padding:8px 14px;cursor:pointer}button:focus-visible{outline:3px solid #58d9b3}aside{position:fixed;bottom:65px;left:5%;right:5%;padding:22px;background:#fff;color:#101d2d;border-radius:8px;font-size:18px;line-height:1.5;max-height:35vh;overflow:auto}aside[hidden]{display:none}@media print{@page{size:16in 9in;margin:0}body{background:white}main{display:block;height:auto}section,section.active{display:block;width:16in;height:9in;break-after:page}nav,aside{display:none}}'''
script='''let current=0;const pages=[...document.querySelectorAll('section')],n=document.querySelector('aside'),count=document.querySelector('#count');function show(i){current=Math.max(0,Math.min(pages.length-1,i));pages.forEach((p,j)=>p.classList.toggle('active',j===current));count.textContent=(current+1)+' / '+pages.length;n.textContent=notes[current];}function toggleNotes(){n.hidden=!n.hidden}document.querySelector('#prev').onclick=()=>show(current-1);document.querySelector('#next').onclick=()=>show(current+1);document.querySelector('#notes').onclick=toggleNotes;document.querySelector('#full').onclick=()=>{if(document.fullscreenElement)document.exitFullscreen();else document.documentElement.requestFullscreen()};document.addEventListener('keydown',e=>{if(['ArrowRight','PageDown',' '].includes(e.key)){e.preventDefault();show(current+1)}if(['ArrowLeft','PageUp'].includes(e.key)){e.preventDefault();show(current-1)}if(e.key==='Home')show(0);if(e.key==='End')show(pages.length-1);if(e.key.toLowerCase()==='n')toggleNotes()});show(0);'''
page='<!doctype html><html lang="pt-BR"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>DSL visual e espacialidade</title><style>'+css+'</style><main>'+''.join('<section>'+v+'</section>' for v in slides)+'</main><nav aria-label="Controles de apresentação"><button id="prev">← Anterior</button><span id="count" aria-live="polite"></span><button id="next">Próximo →</button><button id="notes">Notas (N)</button><button id="full">Tela cheia</button></nav><aside hidden></aside><script>const notes='+json.dumps([n for _,n in notes],ensure_ascii=False)+';'+script+'</script></html>'
(OUT/'apresentacao.html').write_text(page,encoding='utf-8')
(OUT/'roteiro-de-fala.md').write_text('# Roteiro de fala — DSL visual e espacialidade\n\n12 slides · duração sugerida: 10–12 minutos.\n\n'+ '\n\n'.join(f'## {i+1}. {title}\n\n{note}' for i,(title,note) in enumerate(notes))+'\n',encoding='utf-8')
for i,v in enumerate(slides): (OUT/f'slide-{i+1:02}.svg').write_text(v,encoding='utf-8')
print(f'Created {len(slides)} slides in {OUT}')
